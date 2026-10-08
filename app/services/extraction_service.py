import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation as PptxPresentation
from typing import List, Dict, Any

def extract_pdf(file_path: str) -> List[Dict[str, Any]]:
    extracted = []
    doc = fitz.open(file_path)
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text("text")
        if text.strip():
            extracted.append({
                "page_number": str(page_num + 1),
                "text": text.strip(),
                "section_title": None
            })
    return extracted

def extract_docx(file_path: str) -> List[Dict[str, Any]]:
    extracted = []
    doc = DocxDocument(file_path)
    current_section = None
    text_buffer = []
    
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            # flush buffer
            if text_buffer:
                extracted.append({
                    "page_number": None,
                    "text": "\n".join(text_buffer).strip(),
                    "section_title": current_section
                })
                text_buffer = []
            current_section = para.text
        else:
            if para.text.strip():
                text_buffer.append(para.text.strip())
                
    if text_buffer:
        extracted.append({
            "page_number": None,
            "text": "\n".join(text_buffer).strip(),
            "section_title": current_section
        })
    return extracted

def extract_pptx(file_path: str) -> List[Dict[str, Any]]:
    extracted = []
    prs = PptxPresentation(file_path)
    for i, slide in enumerate(prs.slides):
        text_runs = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text_runs.append(run.text.strip())
        
        text = " ".join(text_runs).strip()
        if text:
            extracted.append({
                "page_number": f"Slide {i + 1}",
                "text": text,
                "section_title": None
            })
    return extracted

def extract_txt(file_path: str) -> List[Dict[str, Any]]:
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    if text.strip():
        return [{
            "page_number": None,
            "text": text.strip(),
            "section_title": None
        }]
    return []

def extract_text_from_file(file_path: str, file_type: str) -> List[Dict[str, Any]]:
    ext = file_type.lower()
    if ext == "pdf":
        return extract_pdf(file_path)
    elif ext == "docx":
        return extract_docx(file_path)
    elif ext == "pptx":
        return extract_pptx(file_path)
    elif ext == "txt":
        return extract_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
